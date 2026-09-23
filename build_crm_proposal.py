from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
from datetime import date

OUT = Path(r"E:\crmdemo\deliverables\MANOD_CRM_Enhancement_Proposal.docx")
LOGO = Path(r"E:\crmdemo\CRM\src\assets\manod-logo.jpg")
NAVY = "17365D"; BLUE = "2E75B6"; TEAL = "159A9C"; LIGHT = "EDF3F8"; PALE = "F5F8FA"; GRAY = "667085"; DARK = "202B33"; WHITE = "FFFFFF"; GOLD = "D89B2B"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def margins(cell, top=100, start=130, bottom=100, end=130):
    tcPr = cell._tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for tag, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el = tcMar.find(qn('w:'+tag))
        if el is None: el = OxmlElement('w:'+tag); tcMar.append(el)
        el.set(qn('w:w'), str(value)); el.set(qn('w:type'),'dxa')

def set_cell_text(cell, text, bold=False, color=DARK, size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text); font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

def font(run, name='Aptos', size=11, bold=None, color=DARK, italic=None):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def ptext(doc, text='', size=10.5, bold=False, color=DARK, align=None, before=0, after=7, italic=False, keep=False):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.18
    if align is not None: p.alignment=align
    p.paragraph_format.keep_with_next=keep
    r=p.add_run(text); font(r,size=size,bold=bold,color=color,italic=italic); return p

def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f'Heading {level}')

def bullet(doc, title, detail):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.15
    r=p.add_run(title+': '); font(r,size=10.2,bold=True,color=NAVY)
    r=p.add_run(detail); font(r,size=10.2,color=DARK)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,(h,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Inches(w); set_cell_text(t.rows[0].cells[i],h,True,WHITE,9); shade(t.rows[0].cells[i],NAVY)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,(value,w) in enumerate(zip(row,widths)):
            cells[i].width=Inches(w); set_cell_text(cells[i],str(value),False,DARK,9.1)
            shade(cells[i], WHITE if ridx%2==0 else PALE)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    ptext(doc,'',after=2)
    return t

def callout(doc, label, text):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(6.5)
    c=t.cell(0,0); shade(c,LIGHT); margins(c,160,180,160,180); c.text=''
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3)
    r=p.add_run(label.upper()); font(r,size=8.5,bold=True,color=TEAL)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
    r=p.add_run(text); font(r,size=10.2,bold=True,color=NAVY)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(DARK)
normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.18
for name,size,color,before,after in [('Title',29,NAVY,0,6),('Subtitle',13,GRAY,0,8),('Heading 1',17,NAVY,16,8),('Heading 2',12.5,BLUE,11,5),('Heading 3',11,TEAL,8,4)]:
    s=styles[name]; s.font.name='Aptos Display' if name in ['Title','Heading 1'] else 'Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=name!='Subtitle'
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for lname in ['List Bullet','List Number']:
    s=styles[lname]; s.font.name='Aptos'; s.font.size=Pt(10.2); s.paragraph_format.left_indent=Inches(.38); s.paragraph_format.first_line_indent=Inches(-.19); s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.15

# Header/footer
header=sec.header
ht=header.add_table(rows=1,cols=2,width=Inches(6.5)); ht.autofit=False; ht.columns[0].width=Inches(3.8); ht.columns[1].width=Inches(2.7)
set_cell_text(ht.cell(0,0),'MANOD | CRM SOLUTIONS',True,NAVY,8.5); set_cell_text(ht.cell(0,1),'CLIENT PROPOSAL • CONFIDENTIAL',True,GRAY,8.5); ht.cell(0,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer; fp=footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run('MANOD CRM  |  Enhancement Proposal  |  '); font(r,size=8,color=GRAY)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fp._p.append(field)

# Cover
ptext(doc,'CRM ENHANCEMENT PROPOSAL',size=9,bold=True,color=TEAL,align=WD_ALIGN_PARAGRAPH.CENTER,after=18)
if LOGO.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18); p.add_run().add_picture(str(LOGO),width=Inches(1.3))
p=doc.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Connected Sales, Service & Workforce Automation')
p=doc.add_paragraph(); p.style='Subtitle'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('A phased enhancement of MANOD CRM for lead capture, client communication, field service, intelligent assignment, attendance and payroll')
ptext(doc,'Prepared for: Mr. Arockia Raj, Century Polypacks',size=11,bold=True,color=NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,before=18,after=3)
ptext(doc,f'Prepared by: Hemanshi, MANOD  •  {date.today().strftime("%d %B %Y")}  •  Version 1.1',size=9.5,color=GRAY,align=WD_ALIGN_PARAGRAPH.CENTER,after=20)
callout(doc,'Proposal objective','Create one connected operating system from lead generation to proposal, service delivery, employee attendance and payroll communication.')
heading(doc,'Executive Summary',1)
ptext(doc,'MANOD proposes extending the current CRM into a unified platform for sales, client engagement, machine service and workforce operations. The solution will preserve the existing lead, proposal, follow-up, customer-success and CRM-assistant foundation while adding reliable communication automation, mobile field workflows, map-based lead intelligence, transparent service costing, smart ownership and HRMS automation.')
table(doc,['Business outcome','What improves'],[
('Faster response','Automatic acknowledgements and guided follow-ups through email or WhatsApp.'),
('Better field visibility','Mobile lead capture, GPS location and map access for each company or lead.'),
('Controlled service cost','Separate admin costing and client-facing estimates or invoices.'),
('Clear accountability','CRM assistant recommendations and named ownership for every lead and task.'),
('Accurate workforce records','Clock-in/out, calculated attendance and payroll delivery by email.')],[1.65,4.85])

heading(doc,'1. Current CRM Foundation',1)
ptext(doc,'The proposed work builds on the current MANOD CRM application. The existing codebase includes lead capture and import, duplicate detection, GPS fields, lead assignment, proposals, follow-ups, payment reminders, customer-success journeys, machine/service records, email support and a CRM assistant that can summarize workload and draft messages.')
callout(doc,'Positioning for the client','This proposal combines available CRM capabilities with the requested enhancements. Final feature availability will follow the agreed implementation plan and acceptance testing.')
heading(doc,'2. Proposed Solution Scope',1)
heading(doc,'2.1 Automated client initiation and communication',2)
bullet(doc,'Trigger','When a lead is created, qualified or moved to the Proposal stage, initiate the correct client communication workflow.')
bullet(doc,'Channels','Send by email and/or WhatsApp according to client consent, available contact details and administrator rules.')
bullet(doc,'Content','Share the proposal, quotation, product information, appointment details and documents using approved templates.')
bullet(doc,'Tracking','Record sent time, channel, delivery status, responsible person and next follow-up inside the CRM timeline.')
bullet(doc,'Controls','Allow preview, approval, resend and manual override; keep an audit log of every outbound message.')
heading(doc,'Recommended client initiation message',2)
callout(doc,'Email / WhatsApp template','Dear [Client Name], thank you for your interest in our solutions. Your enquiry has been registered in our CRM and [Assigned Person] will be your point of contact. Proposal details, quotations, documents and service updates will be shared through your registered email address and/or WhatsApp number. Please confirm your preferred communication channel. Regards, MANOD Team.')
heading(doc,'2.2 Mobile lead-generation application',2)
bullet(doc,'Field capture','Create leads from a mobile-friendly application with contact, company, requirement, machine and budget details.')
bullet(doc,'Fast intake','Support camera attachments, business-card/manual entry and optional campaign or QR-code forms.')
bullet(doc,'Offline readiness','Queue entries when connectivity is weak and sync after reconnection, subject to final technical design.')
bullet(doc,'Duplicate prevention','Check phone, email, name and company before saving or importing a lead.')
bullet(doc,'Permissions','Apply role-based access for sales executives, managers, service staff and administrators.')
heading(doc,'2.3 Map and company-location intelligence',2)
bullet(doc,'Lead map','Store latitude, longitude, accuracy and captured time for every lead where location permission is granted.')
bullet(doc,'Company location','Display the customer site on a map and provide open-in-maps navigation for field visits.')
bullet(doc,'Territory view','Filter leads by area, owner, stage and visit status; visualize nearby leads for route planning.')
bullet(doc,'Data quality','Permit address-based geocoding, manual pin correction and location verification.')

heading(doc,'2.4 Machine service and cost management',2)
bullet(doc,'Machine profile','Maintain customer, model, serial number, installation site, warranty and service history.')
bullet(doc,'Service job','Create complaints, visits, engineer assignments, checklists, parts/consumables and completion confirmation.')
bullet(doc,'Admin cost view','Show labour, travel, parts, taxes, internal cost, margin and approval information to authorized users only.')
bullet(doc,'Client cost view','Generate a clean estimate, quotation or service report showing only approved client-facing charges.')
bullet(doc,'Communication','Email or WhatsApp the approved document and record customer acknowledgement or acceptance.')
heading(doc,'2.5 CRM assistant and assignment automation',2)
bullet(doc,'Smart assignment','Recommend or assign a responsible person using territory, product expertise, workload, availability or round-robin rules.')
bullet(doc,'Manager control','Allow reassignment, escalation and assignment history with reason and timestamp.')
bullet(doc,'Assistant actions','Surface missed follow-ups, proposal-stage leads, high-value opportunities, duplicate leads, payment risks and workload imbalance.')
bullet(doc,'Human approval','Keep sensitive actions—pricing, final proposal sending and reassignment—subject to configured approval.')
heading(doc,'2.6 HRMS attendance and payroll',2)
bullet(doc,'Clock-in / clock-out','Employees record attendance from web or mobile; optionally capture time, permitted location and device information.')
bullet(doc,'Automatic calculation','Calculate working hours, late arrival, early departure, overtime, breaks, leave and attendance exceptions using company policy.')
bullet(doc,'Approval workflow','Managers review missing punches, corrections, leave and overtime before payroll is finalized.')
bullet(doc,'Payroll preparation','Use approved attendance, salary components, allowances, deductions and statutory settings to calculate payroll.')
bullet(doc,'Payslip delivery','Generate password-protected payslips where required and email them to each employee; retain a delivery audit trail.')
bullet(doc,'Privacy','Restrict salary and personal information by role, encrypt data in transit and apply retention/access policies.')
heading(doc,'3. End-to-End Workflow',1)
table(doc,['Stage','Automated flow','Primary owner'],[
('1. Capture','Mobile/web lead entry → duplicate check → GPS/address verification','Sales executive'),
('2. Qualify','Requirement and machine details → lead score → assigned person','Sales manager / CRM assistant'),
('3. Engage','Initiation message → preferred channel → follow-up task','Assigned salesperson'),
('4. Propose','Quotation/proposal approval → email/WhatsApp sharing → status tracking','Sales + approver'),
('5. Deliver','Won lead → customer-success plan → machine installation/service record','Operations / service'),
('6. Service','Complaint/visit → estimate → approval → work report → client confirmation','Service team'),
('7. Workforce','Clock-in/out → exceptions approval → payroll → payslip email','Employee / HR / payroll')],[.78,4.35,1.37])

heading(doc,'4. Delivery Plan',1)
table(doc,['Phase','Indicative scope','Estimated duration'],[
('Phase 1','Discovery, workflow confirmation, roles, templates and solution design','Week 1'),
('Phase 2','Email/WhatsApp automation, consent, logs, assignment rules and CRM-assistant enhancements','Weeks 1–2'),
('Phase 3','Mobile lead capture, GPS/maps, territory and field workflows','Weeks 2–3'),
('Phase 4','Machine service jobs, admin/client costing, approvals and documents','Weeks 3–4'),
('Phase 5','HRMS attendance, policy engine, payroll calculation and payslip email','Weeks 4–5'),
('Phase 6','UAT, security review, migration, training and production rollout','Weeks 5–6')],[.75,4.55,1.2])
ptext(doc,'Committed delivery period: 4–6 weeks from proposal acceptance, receipt of the initial payment and timely provision of all required client inputs, credentials and approvals.',size=9.5,bold=True,color=NAVY)
heading(doc,'5. Deliverables',1)
bullet(doc,'Configured CRM','Updated web CRM modules, workflows, roles and dashboards.')
bullet(doc,'Mobile experience','Responsive/PWA or native application option, selected during discovery.')
bullet(doc,'Integrations','Approved email provider, WhatsApp Business Platform/provider and map service integration.')
bullet(doc,'Documents','Proposal, quotation, service estimate/report and payslip templates.')
bullet(doc,'Quality','Functional testing, role/security tests, integration tests and user acceptance support.')
bullet(doc,'Enablement','Administrator guide, user training and go-live handover.')
heading(doc,'6. Assumptions and Dependencies',1)
bullet(doc,'Client inputs','The client will provide branding, users/roles, attendance rules, salary structures, statutory requirements, message templates and approval authorities.')
bullet(doc,'Third-party integrations','All third-party integrations are the client’s responsibility. Century Polypacks will procure and maintain the required accounts, subscriptions, approvals, API access, credentials and provider payments for WhatsApp Business, email, maps/geocoding, hosting, SMS, app stores and any other external service.')
bullet(doc,'Consent and compliance','The client is responsible for lawful contact consent, message content, employee notices and local payroll/statutory compliance; MANOD will configure the agreed controls.')
bullet(doc,'Data migration','Scope depends on the quality, format and volume of available lead, machine, employee and payroll data.')
bullet(doc,'Integration limits','Final feasibility and delivery dates depend on third-party APIs, verification, rate limits and approval timelines.')

heading(doc,'7. Acceptance Criteria',1)
table(doc,['Area','Minimum acceptance outcome'],[
('Communication','Approved trigger sends the correct email/WhatsApp template, records status and creates the next action.'),
('Mobile leads','Authorized user creates a lead with required fields; duplicate and location checks operate as agreed.'),
('Maps','A lead/company can be viewed on the map and opened for navigation with correct permission handling.'),
('Service cost','Admin sees internal cost/margin; client document contains only approved client-facing charges.'),
('Assignment','Rule or assistant assigns/recommends an owner and records manual changes.'),
('Attendance','Clock events and approved exceptions produce correct hours under the signed-off policy test cases.'),
('Payroll','Approved payroll generates the expected payslip and sends it to the employee email with an audit record.')],[1.35,5.15])
heading(doc,'8. Commercial Proposal',1)
callout(doc,'Total quotation','INR 1,00,000 (Indian Rupees One Lakh Only) for the scope described in this proposal.')
doc.add_page_break(); heading(doc,'Payment milestones',2)
table(doc,['Installment','Milestone','Amount'],[
('50%','On proposal acceptance and project commencement','INR 50,000'),
('40%','On completion of the agreed modules and delivery for user acceptance testing','INR 40,000'),
('10%','On final acceptance and production handover','INR 10,000')],[.85,4.15,1.5])
ptext(doc,'All third-party integration accounts, subscriptions, licences, usage fees, approvals, credentials and related provider charges will be arranged and paid directly by Century Polypacks.',size=9.5,bold=True,color=NAVY)
heading(doc,'Scope confirmation and change requests',2)
callout(doc,'Important acceptance condition','Century Polypacks is requested to read this proposal carefully and confirm that the stated scope meets its requirements before signing. Once the proposal is accepted, any addition, alteration or change outside the accepted scope will be evaluated separately and may result in additional charges and a revised delivery schedule. Work on such a change will begin only after written approval of the applicable cost and timeline.')
heading(doc,'9. Next Steps',1)
for text in ['Confirm business workflows, user roles and priority modules.','Select WhatsApp/email/map providers and mobile delivery option.','Conduct HR/payroll policy and statutory requirements workshop.','Approve scope, milestones, commercials and acceptance plan.','Begin Phase 1 discovery and prepare the implementation backlog.']:
    p=doc.add_paragraph(style='List Number'); p.add_run(text)
heading(doc,'Approval',1)
table(doc,['For Century Polypacks','For MANOD'],[
('Name: Mr. Arockia Raj','Name: Rekha Malar'),('Title: _______________________________','Title: Managing Director'),('Signature: ___________________________','Signature: ___________________________'),('Date: _______________________________',f'Date: {date.today().strftime("%d %B %Y")}')],[3.25,3.25])

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.core_properties.title='MANOD CRM Enhancement Proposal'; doc.core_properties.subject='CRM, communication, field service, HRMS and payroll automation proposal'; doc.core_properties.author='MANOD'
doc.save(OUT)
print(OUT)
